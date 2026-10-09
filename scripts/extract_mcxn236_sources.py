#!/usr/bin/env python3
"""Extract the collected FRDM-MCXN236 Rev C BOM and Cadence PST sources.

Read-only source processing. Requires openpyxl==3.1.5; creates JSON/CSV, not an
Excel workbook. Native text grammar is checked in full and unexpected tokens
stop extraction instead of silently dropping records.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from dataclasses import dataclass
import hashlib
import io
import json
from pathlib import Path
import re

import openpyxl


SOURCE_DIR = Path("original-design-files/Design Files")
XLSX_NAME = "SCH-90828_C.BOM.xlsx"
BOM_NAME = "SCH-90828_C.BOM"
EXPECTED_COLUMNS = ["Item", "Quantity", "AgilePN", "Description", "MFG", "MFG_PN", "VALUE", "ASSY_OPT", "FIND#", "PRE_APN", "Reference"]
LEXER = re.compile(r"(?P<space>\s+)|(?P<comment>\{[^}]*\})|(?P<quoted>'[^']*')|(?P<punct>[:;,=])|(?P<word>[^\s'{}:;,=]+)")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


@dataclass
class Token:
    value: str
    kind: str
    line: int


class Parser:
    def __init__(self, text: str, name: str):
        self.name, self.tokens, self.comments = name, [], []
        offset, line = 0, 1
        for match in LEXER.finditer(text):
            require(match.start() == offset, f"{name}: unexpected syntax at line {line}")
            raw, kind = match.group(), match.lastgroup
            if kind == "comment":
                self.comments.append({"line": line, "text": raw})
            elif kind != "space":
                self.tokens.append(Token(raw[1:-1] if kind == "quoted" else raw, kind, line))
            line += raw.count("\n")
            offset = match.end()
        require(offset == len(text), f"{name}: unparsed trailing text")
        self.position, self.last_line = 0, 0

    def peek(self) -> str | None:
        return self.tokens[self.position].value if self.position < len(self.tokens) else None

    def pop(self, expected: str | None = None, kind: str | None = None) -> Token:
        require(self.position < len(self.tokens), f"{self.name}: unexpected EOF")
        token = self.tokens[self.position]
        require(expected is None or token.value == expected, f"{self.name}:{token.line}: expected {expected!r}, got {token.value!r}")
        require(kind is None or token.kind == kind, f"{self.name}:{token.line}: expected {kind}, got {token.kind}")
        self.position += 1
        self.last_line = token.line
        return token

    def properties(self) -> dict:
        """A comma-separated property list with one terminating semicolon."""
        result = {}
        if self.peek() == ";":
            self.pop(";")
            return result
        while True:
            key = self.pop(kind="word").value
            self.pop("=")
            value = self.pop(kind="quoted").value
            require(key not in result, f"{self.name}: duplicate property {key}")
            result[key] = value
            separator = self.pop().value
            require(separator in (",", ";"), f"{self.name}: unexpected property separator {separator}")
            if separator == ";":
                return result

    def body(self, end: str) -> dict:
        result = {}
        while self.peek() != end:
            values = self.properties()
            require(not (set(result) & set(values)), f"{self.name}: duplicate body properties")
            result.update(values)
        self.pop(end)
        self.pop(";")
        return result

    def header(self, expected_type: str) -> None:
        self.pop("FILE_TYPE")
        self.pop("=")
        self.pop(expected_type)
        self.pop(";")

    def finish(self) -> None:
        self.pop("END.")
        require(self.peek() is None, f"{self.name}: trailing tokens after END.")


def parse_parts(text: str) -> dict:
    parser = Parser(text, "pstxprt.dat")
    parser.header("EXPANDEDPARTLIST")
    parser.pop("DIRECTIVES")
    directives = parser.body("END_DIRECTIVES")
    parts = []
    while parser.peek() == "PART_NAME":
        start = parser.pop("PART_NAME").line
        refdes = parser.pop(kind="word").value
        primitive = parser.pop(kind="quoted").value
        parser.pop(":")
        properties = parser.properties()
        sections = []
        while parser.peek() == "SECTION_NUMBER":
            section_start = parser.pop("SECTION_NUMBER").line
            number = int(parser.pop(kind="word").value)
            instance = parser.pop(kind="quoted").value
            parser.pop(":")
            props = parser.properties()
            page = re.search(r":page([0-9]+)_", props.get("P_PATH", ""))
            sections.append({"number": number, "instance_path": instance, "properties": props,
                             "page_from_P_PATH": int(page.group(1)) if page else None,
                             "source_lines": [section_start, parser.last_line]})
        parts.append({"refdes": refdes, "primitive": primitive, "properties": properties,
                      "sections": sections, "source_lines": [start, parser.last_line]})
    parser.finish()
    require(len(parts) == len(re.findall(r"(?m)^PART_NAME\s*$", text)), "Part marker count differs")
    require(len({part["refdes"] for part in parts}) == len(parts), "Duplicate native part reference")
    return {"comments": parser.comments, "directives": directives, "parts": parts}


def parse_primitives(text: str) -> dict:
    parser = Parser(text, "pstchip.dat")
    parser.header("LIBRARY_PARTS")
    primitives = []
    while parser.peek() == "primitive":
        start = parser.pop("primitive").line
        name = parser.pop(kind="quoted").value
        parser.pop(";")
        parser.pop("pin")
        pins = []
        while parser.peek() != "end_pin":
            pin_token = parser.pop(kind="quoted")
            parser.pop(":")
            props = {}
            while parser.peek() != "end_pin" and parser.tokens[parser.position].kind != "quoted":
                added = parser.properties()
                require(not (set(props) & set(added)), "Duplicate pin property")
                props.update(added)
            numbers = re.fullmatch(r"\(([^()]*)\)", props.get("PIN_NUMBER", ""))
            require(numbers is not None, f"Unsupported PIN_NUMBER: {props.get('PIN_NUMBER')}")
            pins.append({"symbol_pin_name": pin_token.value, "properties": props,
                         "physical_pins": [pin.strip() for pin in numbers.group(1).split(",")],
                         "source_lines": [pin_token.line, parser.last_line]})
        parser.pop("end_pin")
        parser.pop(";")
        parser.pop("body")
        body = parser.body("end_body")
        parser.pop("end_primitive")
        parser.pop(";")
        primitives.append({"name": name, "body_properties": body, "pins": pins,
                           "source_lines": [start, parser.last_line]})
    parser.finish()
    require(len(primitives) == len(re.findall(r"(?m)^primitive ", text)), "Primitive marker count differs")
    require(len({p["name"] for p in primitives}) == len(primitives), "Duplicate primitive name")
    return {"comments": parser.comments, "primitives": primitives}


def parse_nets(text: str) -> dict:
    parser = Parser(text, "pstxnet.dat")
    parser.header("EXPANDEDNETLIST")
    records = []
    while parser.peek() == "NET_NAME":
        start = parser.pop("NET_NAME").line
        name = parser.pop(kind="quoted").value
        path = parser.pop(kind="quoted").value
        parser.pop(":")
        props = parser.properties()
        nodes = []
        while parser.peek() == "NODE_NAME":
            node_start = parser.pop("NODE_NAME").line
            refdes, pin = parser.pop(kind="word").value, parser.pop(kind="word").value
            instance = parser.pop(kind="quoted").value
            parser.pop(":")
            symbol_pin = parser.pop(kind="quoted").value
            parser.pop(":")
            node_props = parser.properties()
            nodes.append({"refdes": refdes, "pin": pin, "symbol_pin_name": symbol_pin,
                          "instance_path": instance, "properties": node_props,
                          "source_lines": [node_start, parser.last_line]})
        no_connect = name == "NC" and path == "NC" and props.get("C_SIGNAL") == "NC"
        records.append({"id": f"net-record-{len(records) + 1:03d}", "name": name,
                        "instance_path": path, "properties": props,
                        "record_kind": "no_connect_group" if no_connect else "electrical_net",
                        "nodes": nodes, "source_lines": [start, parser.last_line]})
    parser.finish()
    require(len(records) == len(re.findall(r"(?m)^NET_NAME\s*$", text)), "Net marker count differs")
    require(sum(len(n["nodes"]) for n in records) == len(re.findall(r"(?m)^NODE_NAME\s", text)), "Node marker count differs")
    require(len({n["name"] for n in records}) == len(records), "Duplicate net record name")
    return {"comments": parser.comments,
            "no_connect_policy": "A record with name, instance_path and C_SIGNAL all equal to NC is preserved as a no-connect group. Its pins must not be treated as electrically joined.",
            "records": records}


def parse_references(value: str) -> list[str]:
    refs = [ref.strip() for ref in value.split(",")]
    require(all(re.fullmatch(r"[A-Z]+[1-9][0-9]*", ref) for ref in refs), f"Unrecognized reference syntax: {value!r}")
    require(len(refs) == len(set(refs)), f"Duplicate reference within BOM row: {value!r}")
    return refs


def extract_workbook(path: Path) -> tuple[dict, list[dict]]:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    sheets, items = [], []
    for sheet in workbook:
        header_row = None
        raw_rows = []
        for row in sheet.iter_rows():
            values = [cell.value for cell in row]
            require(not any(cell.data_type == "f" for cell in row), "BOM formulas require explicit formula/cache handling")
            raw_rows.append({"row": row[0].row, "values": values,
                             "cell_types": [cell.data_type for cell in row],
                             "number_formats": [cell.number_format for cell in row]})
            if values == EXPECTED_COLUMNS:
                require(header_row is None, "Multiple BOM header rows need explicit handling")
                header_row = row[0].row
            elif header_row is not None and isinstance(values[0], (int, float)):
                require(isinstance(values[0], int), "BOM item number is not an integer")
                raw = dict(zip(EXPECTED_COLUMNS, values))
                refs = parse_references(raw["Reference"])
                require(raw["Quantity"] == len(refs), f"BOM quantity/reference mismatch at row {row[0].row}")
                item = {"id": f"bom-item-{raw['Item']:03d}", "sheet": sheet.title,
                        "row": row[0].row, "range": f"A{row[0].row}:K{row[0].row}",
                        "source_columns": raw, "reference_designators": refs,
                        "dnp_explicit": str(raw["ASSY_OPT"]).strip().upper() == "DNP",
                        "manufacturer_part_number_text": None if raw["MFG_PN"] is None else str(raw["MFG_PN"]),
                        "manufacturer_part_number_cell_type": row[5].data_type}
                items.append(item)
            elif header_row is not None and any(value is not None for value in values):
                require(all(value is None or str(value).strip("_") == "" for value in values), f"Unparsed BOM content at row {row[0].row}")
        require(header_row is not None, f"BOM headers missing from {sheet.title}")
        sheets.append({"sheet": sheet.title, "max_row": sheet.max_row,
                       "max_column": sheet.max_column, "header_row": header_row,
                       "columns": EXPECTED_COLUMNS, "rows": raw_rows})
    workbook.close()
    refs = [ref for item in items for ref in item["reference_designators"]]
    require(len(refs) == len(set(refs)), "Reference appears in multiple BOM items")
    require(len({item["id"] for item in items}) == len(items), "Duplicate BOM item number")
    return {"library": "openpyxl", "library_version": openpyxl.__version__, "sheets": sheets}, items


def parse_text_bom(text: str) -> dict:
    lines = text.splitlines()
    headers = [index for index, line in enumerate(lines) if line.startswith("Item\tQuantity\t")]
    require(len(headers) == 1, "Text BOM should have one header")
    header_index = headers[0]
    columns = lines[header_index].split("\t")
    items = []
    for index, line in enumerate(lines[header_index + 1:], header_index + 2):
        if not line.strip() or line.startswith("_"):
            continue
        values = line.split("\t")
        require(len(values) == len(columns), f"Text BOM field count at line {index}")
        if values[0].isdigit():
            items.append({"source_lines": [index, index], "raw_lines": [line], "source_columns": dict(zip(columns, values))})
        else:
            require(items and all(value == "" for value in values[:-1]), f"Unexpected BOM continuation line {index}")
            items[-1]["source_columns"]["Reference"] += values[-1]
            items[-1]["raw_lines"].append(line)
            items[-1]["source_lines"][1] = index
    for item in items:
        item["reference_designators"] = parse_references(item["source_columns"]["Reference"])
        require(int(item["source_columns"]["Quantity"]) == len(item["reference_designators"]), "Text BOM quantity mismatch")
    return {"preamble_lines": lines[:header_index], "header_line": header_index + 1,
            "columns": columns, "items": items}


def source_record(path: Path, bundle: Path) -> dict:
    return {"path": path.relative_to(bundle).as_posix(), "sha256": sha256(path), "bytes": path.stat().st_size}


def normalized_field(value: object) -> str:
    return "" if value is None else str(value).strip()


def csv_text(rows: list[list]) -> str:
    stream = io.StringIO(newline="")
    csv.writer(stream, lineterminator="\n").writerows(rows)
    return stream.getvalue()


def build(bundle: Path) -> tuple[dict[str, object], dict[str, str]]:
    source = bundle / SOURCE_DIR
    sources = {path.relative_to(source).as_posix(): source_record(path, bundle) for path in sorted(source.rglob("*")) if path.is_file()}
    workbook, items = extract_workbook(source / XLSX_NAME)
    text_bom = parse_text_bom((source / BOM_NAME).read_text(encoding="utf-8-sig"))
    parts = parse_parts((source / "allegro/pstxprt.dat").read_text(encoding="utf-8-sig"))
    primitives = parse_primitives((source / "allegro/pstchip.dat").read_text(encoding="utf-8-sig"))
    nets = parse_nets((source / "allegro/pstxnet.dat").read_text(encoding="utf-8-sig"))
    index_path = bundle / "schematic-reference-index.json"
    schematic_index = json.loads(index_path.read_text(encoding="utf-8"))
    indexed_pdf = (index_path.parent / schematic_index["source"]["pdf_path_relative_to_index"]).resolve()
    require(sha256(indexed_pdf) == schematic_index["source"]["pdf_sha256"], "Schematic index PDF hash is stale")
    schematic_sources = {"index": source_record(index_path, bundle), "pdf": source_record(indexed_pdf, bundle),
                         "match_basis": "Same normalized reference text only; symbol identity and photo identity are not established by this join."}
    occurrences = {occurrence["id"]: occurrence for occurrence in schematic_index["occurrences"]}
    by_ref = {part["refdes"]: part for part in parts["parts"]}
    by_primitive = {primitive["name"]: primitive for primitive in primitives["primitives"]}
    bom_by_ref = {ref: item for item in items for ref in item["reference_designators"]}
    text_by_ref = {ref: item for item in text_bom["items"] for ref in item["reference_designators"]}
    require(len(text_by_ref) == sum(len(item["reference_designators"]) for item in text_bom["items"]), "Duplicate refdes in text BOM")
    node_refs = {node["refdes"] for net in nets["records"] for node in net["nodes"]}
    all_bom_refs, all_part_refs = set(bom_by_ref), set(by_ref)
    differences = {
        "xlsx_refs_missing_from_text_bom": sorted(all_bom_refs - set(text_by_ref)),
        "text_bom_refs_missing_from_xlsx": sorted(set(text_by_ref) - all_bom_refs),
        "bom_refs_missing_from_pstxprt": sorted(all_bom_refs - all_part_refs),
        "pstxprt_refs_missing_from_bom": sorted(all_part_refs - all_bom_refs),
        "pstxprt_refs_missing_from_pstxnet": sorted(all_part_refs - node_refs),
        "pstxnet_refs_missing_from_pstxprt": sorted(node_refs - all_part_refs),
        "missing_primitive_definitions": sorted({part["primitive"] for part in parts["parts"]} - set(by_primitive)),
        "unused_primitive_definitions": sorted(set(by_primitive) - {part["primitive"] for part in parts["parts"]}),
        "bom_text_common_field_differences": [], "bom_native_part_number_differences": [],
        "dnp_flag_differences": [], "node_pin_not_defined_in_primitive": [],
        "node_instance_not_in_component_sections": [], "duplicate_pin_assignments": [],
    }
    for ref in sorted(all_bom_refs & set(text_by_ref)):
        left, right = bom_by_ref[ref]["source_columns"], text_by_ref[ref]["source_columns"]
        for column in text_bom["columns"]:
            if normalized_field(left[column]) != normalized_field(right[column]):
                difference = {"bom_item": left["Item"], "column": column, "xlsx_value": left[column], "text_value": right[column]}
                if difference not in differences["bom_text_common_field_differences"]:
                    differences["bom_text_common_field_differences"].append(difference)
    connections: dict[str, list] = defaultdict(list)
    pin_assignments: dict[tuple[str, str], list] = defaultdict(list)
    electrical_refs = set()
    for net in nets["records"]:
        for node in net["nodes"]:
            ref, pin = node["refdes"], node["pin"]
            connections[ref].append({"net_record_id": net["id"], "net_name": net["name"],
                                     "record_kind": net["record_kind"], "pin": pin,
                                     "symbol_pin_name": node["symbol_pin_name"], "source_lines": node["source_lines"]})
            pin_assignments[(ref, pin)].append(net["id"])
            if net["record_kind"] == "electrical_net":
                electrical_refs.add(ref)
            if ref in by_ref and by_ref[ref]["primitive"] in by_primitive:
                primitive = by_primitive[by_ref[ref]["primitive"]]
                physical_pins = {p for entry in primitive["pins"] for p in entry["physical_pins"]}
                if pin not in physical_pins:
                    differences["node_pin_not_defined_in_primitive"].append({"refdes": ref, "pin": pin, "net": net["name"]})
                if node["instance_path"] not in {s["instance_path"] for s in by_ref[ref]["sections"]}:
                    differences["node_instance_not_in_component_sections"].append({"refdes": ref, "pin": pin, "instance_path": node["instance_path"]})
    for (ref, pin), assignments in sorted(pin_assignments.items()):
        if len(assignments) > 1:
            differences["duplicate_pin_assignments"].append({"refdes": ref, "pin": pin, "net_record_ids": assignments})
    components, value_text_differences = [], []
    for ref in sorted(all_bom_refs | all_part_refs):
        item, part = bom_by_ref.get(ref), by_ref.get(ref)
        primitive = by_primitive.get(part["primitive"]) if part else None
        body = primitive["body_properties"] if primitive else {}
        if item and part:
            raw = item["source_columns"]
            if normalized_field(raw["AgilePN"]) != normalized_field(body.get("PART_NUMBER")):
                differences["bom_native_part_number_differences"].append({"refdes": ref, "AgilePN": raw["AgilePN"], "PART_NUMBER": body.get("PART_NUMBER")})
            flags = [part["properties"].get("BOM_IGNORE")] + [section["properties"].get("BOM_IGNORE") for section in part["sections"]]
            if any((str(flag).strip().upper() == "DNP") != item["dnp_explicit"] for flag in flags):
                differences["dnp_flag_differences"].append({"refdes": ref, "ASSY_OPT": raw["ASSY_OPT"], "native_BOM_IGNORE": flags})
            if normalized_field(raw["VALUE"]) != normalized_field(body.get("VALUE")):
                value_text_differences.append({"refdes": ref, "bom_VALUE": raw["VALUE"], "primitive_VALUE": body.get("VALUE")})
        components.append({"refdes": ref, "bom_item_id": item["id"] if item else None,
                           "bom_source": {"file": sources[XLSX_NAME]["path"], "sheet": item["sheet"], "row": item["row"], "range": item["range"]} if item else None,
                           "bom_columns": item["source_columns"] if item else None,
                           "manufacturer_part_number_text": item["manufacturer_part_number_text"] if item else None,
                           "dnp_explicit_in_bom": item["dnp_explicit"] if item else None,
                           "native_part": part, "primitive_body_properties": body,
                           "connections": connections[ref],
                           "schematic_text_candidates": [
                               {key: occurrences[identifier][key] for key in ("id", "page", "bbox_pt", "raw_text", "context")}
                               for identifier in schematic_index["by_refdes"].get(ref, [])]})
    no_connect = [net for net in nets["records"] if net["record_kind"] == "no_connect_group"]
    unicode_replacement_cells = []
    for sheet in workbook["sheets"]:
        for row in sheet["rows"]:
            for column, value in enumerate(row["values"], 1):
                if isinstance(value, str) and "\ufffd" in value:
                    unicode_replacement_cells.append(f"{sheet['sheet']}!{openpyxl.utils.get_column_letter(column)}{row['row']}")
    stats = {
        "bom_items": len(items), "bom_reference_designators": len(all_bom_refs),
        "bom_declared_quantity_sum": sum(item["source_columns"]["Quantity"] for item in items),
        "bom_items_marked_DNP": sum(item["dnp_explicit"] for item in items),
        "bom_references_marked_DNP": sum(len(item["reference_designators"]) for item in items if item["dnp_explicit"]),
        "no_part_to_order_references": sum(len(item["reference_designators"]) for item in items if "NO PART TO ORDER" in str(item["source_columns"]["Description"]).upper()),
        "native_parts": len(parts["parts"]), "primitive_definitions": len(primitives["primitives"]),
        "net_records_including_NC": len(nets["records"]), "electrical_net_records": len(nets["records"]) - len(no_connect),
        "node_records_including_NC": sum(len(net["nodes"]) for net in nets["records"]),
        "no_connect_group_nodes": sum(len(net["nodes"]) for net in no_connect),
        "electrical_net_nodes": sum(len(net["nodes"]) for net in nets["records"] if net["record_kind"] == "electrical_net"),
        "references_on_electrical_nets": len(electrical_refs),
        "components_with_schematic_text_candidates": sum(bool(component["schematic_text_candidates"]) for component in components),
    }
    report = {
        "schema_version": "1.0", "board": "FRDM-MCXN236", "source_revision": "C",
        "source_files": sources, "schematic_join_sources": schematic_sources, "statistics": stats,
        "validation": {"native_grammar_fully_consumed": True, "record_markers_match_parsed_counts": True,
                       "bom_quantities_match_reference_lists": True, "cross_source_discrepancies": differences,
                       "cross_source_comparisons_passed": all(not value for value in differences.values())},
        "preserved_source_observations": {
            "value_text_differences": value_text_differences,
            "xlsx_cells_containing_unicode_replacement_character": unicode_replacement_cells,
            "references_only_in_no_connect_group": sorted(node_refs - electrical_refs),
            "references_without_schematic_text_candidates": [component["refdes"] for component in components if not component["schematic_text_candidates"]],
            "assembly_options_raw": [{"value": option, "bom_item_count": count} for option, count in Counter(item["source_columns"]["ASSY_OPT"] for item in items).items()],
        },
        "limitations": [
            "These are Rev C design records, not observations of what is fitted to the photographed specimen; the specimen revision remains unconfirmed.",
            "BOM quantity includes DNP rows, mounting holes, test pads and solder jumpers; it is not a count of parts to buy or parts visible in a photo.",
            "Raw XLSX columns retain original typed values. Numeric MFG_PN cells also have an explicit text identifier for lookup; no digits are added.",
            "ASSY_OPT and PRE_APN are preserved literally. No alternate manufacturer-part list is supplied in these BOM columns, and PRE_APN is not converted into one.",
            "NC is preserved as a no-connect group and must not be used as an electrical short between all listed pins.",
            "Primitive HEIGHT strings have no added units; JEDEC_TYPE is retained as the native footprint identifier, without inferring geometry.",
            "This script parses BOM/XLSX and PST text only. It does not decode BRD/DSN geometry or create Gerbers, STEP, GLB or photo bounding boxes.",
            "Schematic candidates match normalized text only. BGA ball names, notes and old references can share that text; a matching candidate is not a confirmed symbol instance.",
        ],
    }
    wrapped = lambda key, payload: {"schema_version": "1.0", "source": sources[key], **payload}
    json_files = {
        "bom-workbook.json": wrapped(XLSX_NAME, workbook),
        "bom-items.json": wrapped(XLSX_NAME, {"columns": EXPECTED_COLUMNS, "items": items}),
        "bom-text.json": wrapped(BOM_NAME, text_bom),
        "native-parts.json": wrapped("allegro/pstxprt.dat", parts),
        "primitives.json": wrapped("allegro/pstchip.dat", primitives),
        "netlist.json": wrapped("allegro/pstxnet.dat", nets),
        "components.json": {"schema_version": "1.0", "sources": sources, "schematic_join_sources": schematic_sources, "components": components},
        "extraction-report.json": report,
    }
    csv_files = {
        "bom-SCH-90828_C.csv": csv_text([EXPECTED_COLUMNS] + [[item["source_columns"][column] for column in EXPECTED_COLUMNS] for item in items]),
        "net-pins.csv": csv_text([["net_record_id", "net", "record_kind", "refdes", "pin", "symbol_pin_name", "source_line"]] +
                                  [[net["id"], net["name"], net["record_kind"], node["refdes"], node["pin"], node["symbol_pin_name"], node["source_lines"][0]] for net in nets["records"] for node in net["nodes"]]),
    }
    return json_files, csv_files


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--bundle", type=Path, default=Path("data/board-sources/nxp-frdm-mcxn236"))
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--check", action="store_true", help="Regenerate in memory and compare every saved JSON/CSV")
    args = parser.parse_args()
    bundle = args.bundle.resolve()
    output = args.output_dir.resolve() if args.output_dir else bundle / "structured-rev-c"
    json_files, csv_files = build(bundle)
    if args.check:
        for name, value in json_files.items():
            require(json.loads((output / name).read_text(encoding="utf-8")) == value, f"Saved JSON differs from source extraction: {name}")
        for name, value in csv_files.items():
            require((output / name).read_text(encoding="utf-8") == value, f"Saved CSV differs from source extraction: {name}")
    else:
        output.mkdir(parents=True, exist_ok=True)
        for name, value in json_files.items():
            (output / name).write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")
        for name, value in csv_files.items():
            (output / name).write_text(value, encoding="utf-8", newline="\n")
    report = json_files["extraction-report.json"]
    print(json.dumps({"operation": "validated" if args.check else "extracted", "output": str(output),
                      "statistics": report["statistics"], "validation": report["validation"]}, indent=2))


if __name__ == "__main__":
    main()
